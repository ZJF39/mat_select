# -*- coding: utf-8 -*-
"""数据分享闭环：导出（E1）与导入（E2）（契约 §4.4 / §4.5 / §4.6(a)，PRD E1/E2）。

关键约束：
- **checksum 算法冻结**（契约 §4.6(a)）：仅对 materials 数组做 canonical JSON（
  `ensure_ascii=False, sort_keys=True, separators=(",", ":")`）后取 sha256，前缀 `sha256:`。
  导出与导入两侧必须完全一致，否则往返校验会失败。
- 导入分两段：`parse` **只校验不写库**，规范化载荷存入 `import_staging` 并返回 token；
  `commit` 才按冲突策略写库，随后重建 FTS 索引。
"""
from __future__ import annotations

import hashlib
import io
import json
from datetime import datetime, timedelta, timezone
from typing import List, Optional, Tuple

from app.core.errors import import_rejected, not_found, validation_error
from app.core.logging import write_log
from app.db.connection import json_dumps, json_loads, new_uid, now_iso
from app.repository import material_repo, task_repo
from app.repository.base import execute, query_all, query_one, scalar
from app.services import material_service

_TZ = timezone(timedelta(hours=8))
PACK_VERSION = 1

# 导出包内材料必须齐备的必填字段（PRD E2「字段缺失 → 明确报错」）
REQUIRED_FIELDS = ("uid", "name")

XLSX_HEADERS = [
    ("name", "材料名称"), ("short_name", "简称"), ("category_name", "类别"),
    ("grade_type", "材料类型"), ("aliases", "别名"),
    ("density", "密度 g/cm³"), ("tensile_strength", "拉伸强度 MPa"),
    ("elastic_modulus", "弹性模量 GPa"), ("elongation", "断裂伸长率 %"),
    ("notch_impact", "缺口冲击 kJ/m²"), ("hdt", "热变形温度 °C"),
    ("service_temp", "长期使用温度 °C"), ("service_temp_limit", "长期耐温上限 °C"),
    ("price", "参考价"), ("molding_process", "推荐成型工艺"),
    ("features", "主要特性"), ("applications", "典型应用"),
    ("cautions", "注意项"), ("certifications", "认证合规"),
    ("limitations", "失效模式"), ("source", "数据来源"),
]


# ---------------------------------------------------------------- checksum


def pack_checksum(materials: List[dict]) -> str:
    """契约 §4.6(a)：校验对象**仅** materials 数组，checksum 字段自身不参与计算。"""
    canonical = json.dumps(materials, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return "sha256:" + hashlib.sha256(canonical.encode("utf-8")).hexdigest()


# ---------------------------------------------------------------- 取值助手


def _fmt_range(lo, hi) -> str:
    if lo is None and hi is None:
        return ""
    if lo is None:
        return f"≤{hi}"
    if hi is None:
        return str(lo)
    return str(lo) if lo == hi else f"{lo}–{hi}"


def _cell_value(m: dict, key: str) -> str:
    if key in ("density", "tensile_strength", "elastic_modulus", "elongation",
               "notch_impact", "hdt", "service_temp"):
        return _fmt_range(m.get(f"{key}_min"), m.get(f"{key}_max"))
    if key == "price":
        r = _fmt_range(m.get("price_min"), m.get("price_max"))
        unit = m.get("price_unit") or ""
        return f"{r} {unit}".strip()
    if key in ("aliases", "features", "applications", "limitations", "molding_process"):
        v = m.get(key) or []
        return "、".join(str(x) for x in v)
    if key == "cautions":
        v = m.get(key) or []
        return "；".join(
            str(x.get("content")) for x in v if isinstance(x, dict) and x.get("content")
        )
    if key == "certifications":
        v = m.get(key) or {}
        return "；".join(f"{k}={val}" for k, val in v.items() if val not in (None, "", False))
    if key == "service_temp_limit":
        v = m.get(key)
        return "" if v is None else str(v)
    v = m.get(key)
    return "" if v is None else str(v)


# ---------------------------------------------------------------- 导出


def _collect(scope: str, ids: Optional[List[str]], task_id: Optional[int], dbg: Optional[dict]):
    """按范围取材料详情列表。"""
    if scope == "selected":
        uids = ids or []
        out = []
        for uid in uids:
            d = material_repo.get_material(uid)
            if d:
                out.append(d)
        return out

    if scope == "shortlist":
        if not task_id:
            raise validation_error("导出待选清单需指定 task_id")
        task_repo.require(int(task_id))
        out = []
        for item in query_all(
            "SELECT material_id FROM shortlist_item WHERE task_id=? "
            "ORDER BY COALESCE(sort_order, 1000000) ASC, id ASC",
            (int(task_id),),
        ):
            d = material_repo.get_material_by_id(item["material_id"])
            if d:
                out.append(d)
        return out

    dbg = dbg or {}
    if scope == "filtered":
        _, items = material_service.list_materials(
            q=dbg.get("q"), category_ids=dbg.get("category_ids"),
            processes=dbg.get("processes"), temp_min=dbg.get("temp_min"),
            price_max=dbg.get("price_max"), flame=dbg.get("flame"),
            features=dbg.get("features"), sort=dbg.get("sort") or "updated_at",
            order=dbg.get("order") or "desc", page=1, page_size=500,
        )
        return [material_repo.get_material(i["uid"]) for i in items if i.get("uid")]

    # all
    _, items = material_service.list_materials(page=1, page_size=500)
    return [material_repo.get_material(i["uid"]) for i in items if i.get("uid")]


def _work_data(task_id: Optional[int]) -> dict:
    tasks = task_repo.list_tasks(status=None, q=None)
    messages, shortlist, feedback = [], [], []
    for t in tasks:
        messages.extend(
            [dict(x, task_id=t["id"]) for x in query_all(
                "SELECT id, role, content_json, created_at FROM task_message WHERE task_id=? ORDER BY id",
                (t["id"],))]
        )
        shortlist.extend(
            [dict(x, task_id=t["id"]) for x in query_all(
                "SELECT * FROM shortlist_item WHERE task_id=? ORDER BY id", (t["id"],))]
        )
    feedback = [dict(r) for r in query_all("SELECT * FROM recommendation_feedback ORDER BY id")]
    return {"tasks": tasks, "messages": messages, "shortlist": shortlist, "feedback": feedback}


def build_pack(scope: str, include_work_data: bool, ids=None, task_id=None, dbg=None) -> dict:
    materials = _collect(scope, ids, task_id, dbg)
    materials = [m for m in materials if m]
    pack = {
        "pack_version": PACK_VERSION,
        "exported_at": now_iso(),
        "exported_by": "本机用户",
        "material_count": len(materials),
        "checksum": pack_checksum(materials),
        "materials": materials,
    }
    if include_work_data:
        pack["work_data"] = _work_data(task_id)
    return pack


def _filename(fmt: str) -> str:
    d = datetime.now(_TZ).strftime("%Y%m%d")
    return {
        "json": f"MatSelect_材料包_{d}.json",
        "xlsx": f"MatSelect_材料清单_{d}.xlsx",
        "md": f"MatSelect_对比表_{d}.md",
    }.get(fmt, f"MatSelect_导出_{d}.bin")


def _to_xlsx(materials: List[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "材料清单"
    ws.append([cn for _, cn in XLSX_HEADERS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for m in materials:
        ws.append([_cell_value(m, k) for k, _ in XLSX_HEADERS])
    widths = [26, 12, 12, 12, 28, 10, 14, 14, 14, 16, 14, 18, 18, 20, 16, 26, 26, 30, 24, 26, 22]
    for i, w in enumerate(widths[: len(XLSX_HEADERS)], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _to_md(materials: List[dict]) -> str:
    cols = [("name", "材料名称"), ("category_name", "类别"), ("density", "密度 g/cm³"),
            ("tensile_strength", "拉伸强度 MPa"), ("service_temp_limit", "耐温上限 °C"),
            ("price", "参考价"), ("molding_process", "推荐工艺")]
    head = "| " + " | ".join(cn for _, cn in cols) + " |"
    sep = "| " + " | ".join("---" for _ in cols) + " |"
    rows = [
        "| " + " | ".join(_cell_value(m, k) for k, _ in cols) + " |" for m in materials
    ]
    return "\n".join([head, sep] + rows) + "\n"


def export(scope: str, fmt: str, include_work_data: bool, ids=None, task_id=None, dbg=None):
    """返回 (filename, bytes, media_type)。"""
    if scope not in ("all", "filtered", "selected", "shortlist"):
        raise validation_error("scope 仅支持 all / filtered / selected / shortlist")
    if fmt not in ("json", "xlsx", "md"):
        raise validation_error("format 仅支持 json / xlsx / md")

    pack = build_pack(scope, include_work_data, ids, task_id, dbg)
    materials = pack["materials"]
    name = _filename(fmt)

    if fmt == "json":
        body = json.dumps(pack, ensure_ascii=False, indent=2).encode("utf-8")
        media = "application/json"
    elif fmt == "xlsx":
        body = _to_xlsx(materials)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        body = _to_md(materials).encode("utf-8")
        media = "text/markdown"

    write_log("导出", name, {"scope": scope, "format": fmt, "count": len(materials)})
    return name, body, media


# ---------------------------------------------------------------- 导入（两段式）


def _classify(materials: List[dict]):
    """把待导入材料分为 add / update / conflict / invalid。"""
    details = []
    counters = {"add": 0, "update": 0, "conflict": 0, "invalid": 0}
    for m in materials:
        if not isinstance(m, dict):
            details.append({"uid": None, "name": "", "kind": "invalid", "reason": "条目不是对象"})
            counters["invalid"] += 1
            continue
        missing = [f for f in REQUIRED_FIELDS if not str(m.get(f) or "").strip()]
        if missing:
            details.append({
                "uid": m.get("uid"), "name": m.get("name") or "",
                "kind": "invalid", "reason": "缺少必填字段：" + "、".join(missing),
            })
            counters["invalid"] += 1
            continue

        uid = str(m["uid"])
        if scalar("SELECT id FROM material WHERE uid=?", (uid,)) is not None:
            details.append({"uid": uid, "name": m["name"], "kind": "update"})
            counters["update"] += 1
            continue

        name, short = str(m.get("name") or ""), str(m.get("short_name") or "")
        dup = query_one(
            "SELECT uid FROM material WHERE name=? AND COALESCE(short_name,'')=?",
            (name, short),
        )
        if dup:
            details.append({
                "uid": uid, "name": name, "kind": "conflict",
                "reason": f"同名材料已存在（本机 uid={'*' * 6}{str(dup['uid'])[-6:]}）",
            })
            counters["conflict"] += 1
            continue

        details.append({"uid": uid, "name": name, "kind": "add"})
        counters["add"] += 1
    return counters, details


def import_parse(raw: bytes) -> dict:
    """第一段：只校验，不写库；成功则落 import_staging 并返回 token。"""
    try:
        pack = json.loads(raw.decode("utf-8"))
    except Exception as exc:  # noqa: BLE001
        raise import_rejected(f"文件不是合法 JSON：{exc}")

    if not isinstance(pack, dict):
        raise import_rejected("材料包格式不正确：顶层必须是对象")
    version = pack.get("pack_version")
    if version != PACK_VERSION:
        raise import_rejected(f"包版本不兼容：期望 {PACK_VERSION}，实际 {version}")

    materials = pack.get("materials")
    if not isinstance(materials, list):
        raise import_rejected("材料包缺少 materials 数组")

    expected = pack.get("checksum")
    actual = pack_checksum(materials)
    if not expected:
        raise import_rejected("材料包缺少校验和字段，无法验证完整性")
    if expected != actual:
        # PRD E2「失败处理」：校验和不符 → 明确报错并拒绝写入，不产生半截数据。
        # 这里必须在 parse 阶段就拒绝，否则用户可以越过校验直接 commit。
        raise import_rejected(
            "材料包校验和不符（文件可能已损坏或被修改），已拒绝导入且未写入任何数据"
        )

    counters, details = _classify(materials)

    token = new_uid()
    execute(
        "INSERT INTO import_staging(token, payload, created_at) VALUES (?,?,?)",
        (token, json_dumps(pack), now_iso()),
    )

    return {
        "token": token,
        "pack_version": int(version),
        "exported_by": pack.get("exported_by") or "",
        "exported_at": pack.get("exported_at") or "",
        "checksum_ok": True,
        "added": counters["add"],
        "updated": counters["update"],
        "conflicted": counters["conflict"],
        "invalid": counters["invalid"],
        "details": details,
    }


def _insert_material(m: dict, uid: str, name_suffix: str = "") -> None:
    data = {k: v for k, v in m.items() if k not in ("uid", "id", "created_at", "updated_at",
                                                   "category_name", "category_path",
                                                   "negative_feedback", "recent_revisions")}
    if name_suffix:
        data["name"] = f"{data.get('name', '')}{name_suffix}"
    # 分类按 category_path 末位在本地重新解析（跨机器 id 不可信）
    path = m.get("category_path") or []
    if isinstance(path, list) and path:
        row = query_one("SELECT id FROM category WHERE name=?", (str(path[-1]),))
        data["category_id"] = row["id"] if row else None
    material_repo.create_material(data)
    if uid:
        execute("UPDATE material SET uid=? WHERE id=last_insert_rowid()", (uid,))


def import_commit(token: str, conflict_policy: str) -> dict:
    """第二段：按冲突策略写库，随后重建 FTS 索引。"""
    if conflict_policy not in ("skip", "overwrite", "duplicate"):
        raise validation_error("conflict_policy 仅支持 skip / overwrite / duplicate")

    row = query_one("SELECT payload FROM import_staging WHERE token=?", (token,))
    if row is None:
        raise not_found("导入会话不存在或已过期，请重新选择文件")

    pack = json_loads(row["payload"])
    materials = pack.get("materials") or []

    # 纵深防御：commit 时再校验一次暂存载荷的完整性，
    # 防止暂存被篡改或 parse 阶段的历史数据绕过校验。
    if pack.get("checksum") != pack_checksum(materials):
        raise import_rejected("暂存材料包校验和不符，已拒绝写入")

    counters, _ = _classify(materials)

    added = updated = skipped = conflicted = 0
    for m in materials:
        if not isinstance(m, dict) or not str(m.get("name") or "").strip():
            skipped += 1
            continue
        uid = str(m.get("uid") or "")
        existing = scalar("SELECT id FROM material WHERE uid=?", (uid,)) if uid else None

        if existing is None:
            name, short = str(m.get("name")), str(m.get("short_name") or "")
            dup = query_one(
                "SELECT uid, id FROM material WHERE name=? AND COALESCE(short_name,'')=?",
                (name, short),
            )
            if dup is None:
                _insert_material(m, uid)
                added += 1
                continue
            # 冲突：三种策略
            if conflict_policy == "skip":
                skipped += 1
                conflicted += 1
            elif conflict_policy == "overwrite":
                material_repo.update_material(dup["uid"], {
                    k: v for k, v in m.items()
                    if k not in ("uid", "id", "created_at", "updated_at", "category_path")
                })
                updated += 1
            else:  # duplicate：另存副本（更名后新建）
                _insert_material(m, new_uid(), name_suffix="（副本）")
                added += 1
            continue

        if conflict_policy == "skip":
            skipped += 1
        elif conflict_policy == "overwrite":
            material_repo.update_material(uid, {
                k: v for k, v in m.items()
                if k not in ("uid", "id", "created_at", "updated_at", "category_path")
            })
            updated += 1
        else:
            _insert_material(m, new_uid(), name_suffix="（副本）")
            added += 1

    # 重建全文检索索引（导入后必须做，否则新数据搜不到）
    execute("INSERT INTO material_fts(material_fts) VALUES('rebuild')")
    execute("DELETE FROM import_staging WHERE token=?", (token,))

    write_log("导入", f"token={token[:8]}", {"added": added, "updated": updated,
                                            "skipped": skipped, "policy": conflict_policy})
    return {
        "ok": True,
        "added": added,
        "updated": updated,
        "conflicted": conflicted,
        "skipped": skipped,
        "total": added + updated + skipped,
        "message": f"导入完成：新增 {added} 条，更新 {updated} 条，跳过 {skipped} 条",
    }
